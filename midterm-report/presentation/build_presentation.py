from pathlib import Path
from functools import lru_cache
import csv

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "中美外语院系英文简介_5分钟汇报.pptx"
PREVIEW = ROOT / "previews"
PREVIEW.mkdir(parents=True, exist_ok=True)

SLIDE_W, SLIDE_H = 13.333, 7.5
PX_PER_IN = 120
FONT_REG = "/tmp/ppt-fonts/expo/package/400Regular/NotoSansSC_400Regular.ttf"
FONT_BOLD = "/tmp/ppt-fonts/expo/package/700Bold/NotoSansSC_700Bold.ttf"

NAVY = "17324D"
BLUE = "3568A8"
ORANGE = "E38B36"
TEAL = "2B7A78"
TEXT = "23364A"
MUTED = "62748A"
PALE = "F7F9FC"
PALE_BLUE = "EAF1F8"
PALE_ORANGE = "FFF3E5"
PALE_TEAL = "E8F4F2"
WHITE = "FFFFFF"
LINE = "D8E1EB"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def pil_rgb(hex_color):
    return "#" + hex_color


@lru_cache(maxsize=64)
def preview_font(size_px, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size_px)


class Canvas:
    def __init__(self, prs, slide, image):
        self.prs = prs
        self.slide = slide
        self.image = image
        self.draw = ImageDraw.Draw(image)

    @staticmethod
    def px(v):
        return int(round(v * PX_PER_IN))

    def rect(self, x, y, w, h, fill, line=None, radius=0, width=1):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(width)
        else:
            shape.line.fill.background()
        box = [self.px(x), self.px(y), self.px(x + w), self.px(y + h)]
        if radius:
            self.draw.rounded_rectangle(box, radius=self.px(radius), fill=pil_rgb(fill), outline=pil_rgb(line) if line else None, width=max(1, self.px(width / 72)))
        else:
            self.draw.rectangle(box, fill=pil_rgb(fill), outline=pil_rgb(line) if line else None, width=max(1, self.px(width / 72)))
        return shape

    def circle(self, x, y, d, fill, line=None, width=1):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line); shape.line.width = Pt(width)
        else:
            shape.line.fill.background()
        self.draw.ellipse([self.px(x), self.px(y), self.px(x + d), self.px(y + d)], fill=pil_rgb(fill), outline=pil_rgb(line) if line else None)
        return shape

    def line(self, x1, y1, x2, y2, color=LINE, width=1):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
        self.draw.line([self.px(x1), self.px(y1), self.px(x2), self.px(y2)], fill=pil_rgb(color), width=max(1, self.px(width / 72)))
        return shape

    def chevron(self, x, y, w, h, fill):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        points = [
            (self.px(x), self.px(y)),
            (self.px(x + w * 0.55), self.px(y)),
            (self.px(x + w), self.px(y + h / 2)),
            (self.px(x + w * 0.55), self.px(y + h)),
            (self.px(x), self.px(y + h)),
            (self.px(x + w * 0.45), self.px(y + h / 2)),
        ]
        self.draw.polygon(points, fill=pil_rgb(fill))
        return shape

    def text(self, x, y, w, h, text, size=16, color=TEXT, bold=False,
             align="left", valign="top", font="Microsoft YaHei", spacing=1.12):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.space_after = Pt(0); p.space_before = Pt(0)
            for run in p.runs:
                run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
        f = preview_font(max(1, self.px(size / 72)), bold)
        line_h = self.px(size / 72) * spacing
        total_h = line_h * len(lines)
        if valign == "middle":
            ty = self.px(y) + (self.px(h) - total_h) / 2
        elif valign == "bottom":
            ty = self.px(y + h) - total_h
        else:
            ty = self.px(y)
        for line in lines:
            bbox = self.draw.textbbox((0, 0), line, font=f)
            tw = bbox[2] - bbox[0]
            if align == "center":
                tx = self.px(x) + (self.px(w) - tw) / 2
            elif align == "right":
                tx = self.px(x + w) - tw
            else:
                tx = self.px(x)
            self.draw.text((tx, ty), line, font=f, fill=pil_rgb(color))
            ty += line_h
        return box


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.prs.core_properties.title = "基于语料库技术的中美外语院系英文简介话语比较分析"
        self.prs.core_properties.subject = "本科生创新训练培育项目中期汇报"
        self.prs.core_properties.author = "项目组"
        self.slide_no = 0

    def add_slide(self, bg=PALE):
        self.slide_no += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        image = Image.new("RGB", (self.px(SLIDE_W), self.px(SLIDE_H)), pil_rgb(bg))
        canvas = Canvas(self.prs, slide, image)
        canvas.rect(0, 0, SLIDE_W, SLIDE_H, bg)
        return canvas

    @staticmethod
    def px(v):
        return int(round(v * PX_PER_IN))

    def save_slide(self, canvas):
        canvas.image.save(PREVIEW / f"slide_{self.slide_no:02d}.png", quality=95)

    def save(self):
        self.prs.save(OUT)
        print(OUT)


def slide_01(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 1.35, ORANGE)
    c.rect(0.76, 0.52, 3.75, 0.42, NAVY, radius=0.10)
    c.text(0.92, 0.52, 3.43, 0.42, "本科生创新训练培育项目｜中期汇报", 12, WHITE, bold=True, valign="middle")

    c.text(0.78, 1.40, 7.75, 2.12, "基于语料库技术的\n中美外语院系英文简介\n话语比较分析", 30, NAVY, bold=True, spacing=1.28)
    c.text(0.80, 3.76, 7.65, 0.50, "中美高校外语院系英文 About Us 文本的语料库辅助比较", 15, MUTED, valign="middle")
    c.rect(0.80, 4.35, 0.65, 0.045, ORANGE)

    cards = [(0.80, "60", "个院系文本"), (2.86, "30＋30", "中美样本"), (5.15, "10", "核心语步")]
    widths = [1.78, 2.02, 1.78]
    for (x, number, label), w in zip(cards, widths):
        c.rect(x, 4.68, w, 1.03, WHITE, line=LINE, radius=0.10, width=0.8)
        c.text(x, 4.78, w, 0.47, number, 23, BLUE, bold=True, align="center", valign="middle")
        c.text(x, 5.28, w, 0.28, label, 11, MUTED, align="center", valign="middle")

    # Right-side editable visual: a browser card and three research stages.
    c.rect(9.00, 0.95, 3.55, 5.40, PALE_BLUE, radius=0.18)
    c.rect(9.34, 1.36, 2.87, 1.47, WHITE, line=LINE, radius=0.10, width=0.8)
    for i, color in enumerate([ORANGE, BLUE, TEAL]):
        c.circle(9.53 + i * 0.18, 1.55, 0.085, color)
    c.text(9.54, 1.80, 2.46, 0.38, "About Us", 19, NAVY, bold=True)
    c.rect(9.55, 2.28, 2.18, 0.055, LINE, radius=0.02)
    c.rect(9.55, 2.45, 1.62, 0.055, LINE, radius=0.02)

    c.text(9.36, 3.09, 2.80, 0.30, "研究路径", 11, MUTED, bold=True)
    stages = [(3.48, "01", "语步分布", BLUE), (4.20, "02", "语言实现", TEAL), (4.92, "03", "语境解释", ORANGE)]
    for y, num, label, color in stages:
        c.rect(9.34, y, 2.87, 0.56, WHITE, radius=0.09)
        c.circle(9.50, y + 0.11, 0.34, color)
        c.text(9.50, y + 0.11, 0.34, 0.34, num, 9, WHITE, bold=True, align="center", valign="middle")
        c.text(10.02, y + 0.04, 1.88, 0.48, label, 15, NAVY, bold=True, valign="middle")
    c.text(9.36, 5.78, 2.82, 0.28, "MOVE · LANGUAGE · CONTEXT", 8, MUTED, align="center")

    c.line(0.80, 6.30, 12.55, 6.30, LINE, 0.8)
    c.text(0.80, 6.52, 3.00, 0.34, "外语学院", 13, NAVY, bold=True)
    c.text(0.80, 6.93, 6.60, 0.30, "汇报人：________　　指导教师：________", 11, MUTED)
    c.text(10.60, 6.90, 1.20, 0.30, "2026年", 11, MUTED, align="right")
    c.text(11.90, 6.90, 0.65, 0.30, "01 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_02(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 3.20, 0.30, "01　研究定位", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 6.40, 0.52, "研究内容与目标", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "围绕五个相互关联的维度，比较中美外语院系英文简介的机构话语建构",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    c.text(0.80, 2.15, 4.60, 0.32, "研究内容｜五个比较维度", 12, BLUE, bold=True)
    dimensions = [
        ("01", "体裁构成与语步模式", "顺序 · 篇幅 · 出现率"),
        ("02", "教育理念与培养目标", "核心能力 · 工具性—人文性"),
        ("03", "学科形象与学术身份", "评价表达 · 自我指称 · 名词化"),
        ("04", "国际化表述策略", "理想主义 · 工具主义 · 教育主义"),
        ("05", "多语种资源与跨学科整合", "语种资源 · 学科融合 · 社会服务"),
    ]
    for i, (num, title, detail) in enumerate(dimensions):
        y = 2.58 + i * 0.70
        c.rect(0.80, y, 7.36, 0.58, WHITE, line=LINE, radius=0.08, width=0.8)
        c.circle(1.00, y + 0.12, 0.34, BLUE if i < 4 else TEAL)
        c.text(1.00, y + 0.12, 0.34, 0.34, num, 8.5, WHITE, bold=True,
               align="center", valign="middle")
        c.text(1.52, y + 0.07, 3.18, 0.44, title, 12.5, NAVY, bold=True, valign="middle")
        c.text(4.74, y + 0.07, 3.18, 0.44, detail, 10, MUTED,
               align="right", valign="middle")

    c.rect(8.52, 2.15, 3.88, 3.93, PALE_BLUE, radius=0.14)
    c.text(8.80, 2.36, 3.28, 0.34, "研究目标", 12, BLUE, bold=True)
    goals = [
        ("描述", "识别文本结构与语言表达的\n共性及差异", BLUE),
        ("解释", "联系社会文化与制度语境\n阐释差异成因", TEAL),
        ("转化", "形成话语策略建议库\n与研究论文", ORANGE),
    ]
    for i, (label, detail, color) in enumerate(goals):
        y = 2.83 + i * 0.99
        c.rect(8.80, y, 3.32, 0.82, WHITE, radius=0.09)
        c.rect(8.80, y, 0.67, 0.82, color, radius=0.09)
        c.text(8.80, y, 0.67, 0.82, label, 11, WHITE, bold=True,
               align="center", valign="middle")
        c.text(9.67, y + 0.07, 2.20, 0.68, detail, 10.5, TEXT, valign="middle", spacing=1.20)

    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "02 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_03(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 3.20, 0.30, "02　项目概览", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 6.40, 0.52, "项目整体路线", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "定量比较已完成，下一阶段由“描述差异”转向“解释差异与策略转化”",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    # Editable status legend.
    c.circle(10.23, 0.57, 0.12, BLUE)
    c.text(10.42, 0.49, 0.78, 0.28, "已完成", 10, MUTED, valign="middle")
    c.circle(11.33, 0.57, 0.12, ORANGE)
    c.text(11.52, 0.49, 1.02, 0.28, "下一阶段", 10, MUTED, valign="middle")

    stages = [
        ("01", "语料构建", "中国、美国各30个\n外语院系英文简介", "60个院系文本"),
        ("02", "框架与校订", "AI预标注\n人工逐篇校订", "10个核心语步"),
        ("03", "语步分布", "出现率 · 篇幅\n开篇位置", "比较结构组织"),
        ("04", "语言实现", "关键词 · 词丛\n搭配与语义组", "聚焦4个核心语步"),
        ("05", "语境与转化", "跨文化话语分析\n批评话语分析\n社会文化语境解释", "策略建议库 · 论文"),
    ]
    xs = [0.80, 3.18, 5.56, 7.94, 10.32]
    card_y, card_w, card_h = 2.20, 2.08, 3.50
    for i, ((num, title, desc, tag), x) in enumerate(zip(stages, xs)):
        active = i == 4
        accent = ORANGE if active else BLUE
        c.rect(x, card_y, card_w, card_h,
               PALE_ORANGE if active else WHITE,
               line=ORANGE if active else LINE, radius=0.11, width=1.0 if active else 0.8)
        c.rect(x, card_y, card_w, 0.07, accent)
        c.rect(x + 0.18, card_y + 0.25, 0.86 if active else 0.72, 0.30,
               ORANGE if active else PALE_BLUE, radius=0.08)
        c.text(x + 0.18, card_y + 0.25, 0.86 if active else 0.72, 0.30,
               "下一阶段" if active else "已完成", 9,
               WHITE if active else BLUE, bold=True, align="center", valign="middle")
        c.text(x + 0.18, card_y + 0.73, 1.25, 0.54, num, 27, accent, bold=True)
        c.text(x + 0.18, card_y + 1.37, 1.73, 0.42, title, 16, NAVY, bold=True, valign="middle")
        c.text(x + 0.18, card_y + 1.98, 1.73, 0.86, desc, 11, MUTED, spacing=1.25)
        c.line(x + 0.18, card_y + 2.93, x + 1.90, card_y + 2.93,
               ORANGE if active else LINE, 0.8)
        c.text(x + 0.18, card_y + 3.06, 1.73, 0.28, tag, 10,
               ORANGE if active else BLUE, bold=True, valign="middle")

    for i, x in enumerate([2.95, 5.33, 7.71, 10.09]):
        c.chevron(x, 3.75, 0.16, 0.34, ORANGE if i == 3 else LINE)

    c.rect(0.80, 5.98, 11.60, 0.48, PALE_BLUE, radius=0.08)
    c.text(1.00, 5.98, 1.25, 0.48, "本次汇报重点", 10, BLUE, bold=True, valign="middle")
    c.text(2.36, 5.98, 9.70, 0.48,
           "03 语步分布　·　04 语言实现　　｜　　05 语境与转化为下一阶段计划",
           11, TEXT, valign="middle")

    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "03 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_04(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 3.20, 0.30, "03　研究设计", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 8.60, 0.52, "样本、语步与人机协作校订", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "以院系级文本为比较单位，后续统计与语言分析均基于人工校订标签",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    card_y, card_h = 2.20, 3.92
    # Corpus sample.
    c.rect(0.80, card_y, 2.75, card_h, WHITE, line=LINE, radius=0.11, width=0.8)
    c.text(1.04, 2.42, 2.25, 0.32, "01　语料样本", 12, BLUE, bold=True)
    c.text(1.04, 2.88, 2.25, 0.66, "60", 35, NAVY, bold=True, align="center")
    c.text(1.04, 3.50, 2.25, 0.30, "个院系文本", 11, MUTED, align="center")
    c.line(1.04, 3.93, 3.30, 3.93, LINE, 0.8)
    c.rect(1.04, 4.15, 2.26, 0.48, PALE_BLUE, radius=0.08)
    c.circle(1.22, 4.31, 0.14, BLUE)
    c.text(1.50, 4.15, 1.58, 0.48, "中国　30个", 11, NAVY, bold=True, valign="middle")
    c.rect(1.04, 4.77, 2.26, 0.48, PALE_ORANGE, radius=0.08)
    c.circle(1.22, 4.93, 0.14, ORANGE)
    c.text(1.50, 4.77, 1.58, 0.48, "美国　30个", 11, NAVY, bold=True, valign="middle")
    c.text(1.04, 5.45, 2.26, 0.26, "官方网站英文 About Us", 9.5, MUTED, align="center")
    c.text(1.04, 5.75, 2.26, 0.26, "院系级聚合文本", 10, BLUE, bold=True, align="center")

    # Ten-move framework.
    c.rect(3.80, card_y, 5.10, card_h, WHITE, line=LINE, radius=0.11, width=0.8)
    c.text(4.04, 2.42, 1.85, 0.32, "02　语步框架", 12, BLUE, bold=True)
    c.text(5.90, 2.42, 2.74, 0.32, "24个细类 → 10个核心语步", 10, MUTED,
           align="right", valign="middle")
    c.line(4.04, 2.84, 8.66, 2.84, LINE, 0.8)
    moves = [
        ("M01", "机构身份与组织治理"), ("M02", "历史沿革与学脉传承"),
        ("M03", "使命愿景与公共价值"), ("M04", "学科价值与研究取向"),
        ("M05", "课程教学与培养目标"), ("M06", "师资与办学支撑"),
        ("M07", "办学声誉与育人成效"), ("M08", "合作网络与社会服务"),
        ("M09", "受众互动与行动引导"), ("M10", "信息导引与运行通告"),
    ]
    for i, (code, name) in enumerate(moves):
        col = 0 if i < 5 else 1
        row = i if i < 5 else i - 5
        x = 4.04 + col * 2.32
        y = 3.03 + row * 0.51
        c.text(x, y, 0.43, 0.34, code, 9, BLUE if col == 0 else TEAL,
               bold=True, valign="middle")
        c.text(x + 0.46, y, 1.74, 0.34, name, 9.5, TEXT, valign="middle")
        if row < 4:
            c.line(x, y + 0.41, x + 2.18, y + 0.41, LINE, 0.5)
    c.text(4.04, 5.75, 4.62, 0.25, "完整语步定义与语阶见附录", 9, MUTED)

    # Human-in-the-loop annotation workflow.
    c.rect(9.15, card_y, 3.25, card_h, WHITE, line=LINE, radius=0.11, width=0.8)
    c.text(9.39, 2.42, 2.77, 0.32, "03　人机协作校订", 12, BLUE, bold=True)
    workflow = [
        (2.91, "AI预标注", PALE_BLUE, BLUE),
        (3.69, "人工逐篇审阅", WHITE, NAVY),
        (4.47, "立即修改并保存", PALE_ORANGE, ORANGE),
    ]
    for y, label, fill, color in workflow:
        c.rect(9.43, y, 2.69, 0.48, fill, line=LINE if fill == WHITE else None,
               radius=0.08, width=0.8)
        c.text(9.43, y, 2.69, 0.48, label, 11.5, color, bold=True,
               align="center", valign="middle")
    c.text(10.54, 3.39, 0.46, 0.28, "↓", 15, MUTED, align="center")
    c.text(10.54, 4.17, 0.46, 0.28, "↓", 15, MUTED, align="center")
    c.rect(9.43, 5.20, 2.69, 0.68, NAVY, radius=0.09)
    c.text(9.43, 5.27, 2.69, 0.20, "最终分析对象", 8.5, WHITE, align="center")
    c.text(9.43, 5.48, 2.69, 0.30, "人工校订标签", 13, WHITE, bold=True, align="center")

    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "04 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_05(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 4.20, 0.30, "04　研究结果｜语步分布", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 10.60, 0.52, "语步出现率：美国文本的覆盖更稳定", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "中美样本共享10个语步，国别差距主要集中于M09、M08与M03",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    # Main editable paired-bar chart, driven by table 2-2.
    c.rect(0.80, 2.18, 8.15, 4.12, WHITE, line=LINE, radius=0.11, width=0.8)
    c.text(1.03, 2.37, 2.30, 0.30, "文本出现率（%）", 11, BLUE, bold=True)
    c.rect(6.15, 2.43, 0.16, 0.10, BLUE, radius=0.02)
    c.text(6.39, 2.34, 0.62, 0.28, "中国", 9.5, MUTED, valign="middle")
    c.rect(7.16, 2.43, 0.16, 0.10, ORANGE, radius=0.02)
    c.text(7.40, 2.34, 0.72, 0.28, "美国", 9.5, MUTED, valign="middle")

    plot_x, plot_w = 1.55, 6.45
    for tick in [0, 25, 50, 75, 100]:
        x = plot_x + plot_w * tick / 100
        c.line(x, 2.88, x, 6.03, LINE, 0.5)
        c.text(x - 0.22, 2.64, 0.44, 0.20, str(tick), 8, MUTED, align="center")

    data_path = ROOT.parent / "analysis" / "table_2_2_move_occurrence.tsv"
    with data_path.open(encoding="utf-8-sig", newline="") as f:
        occurrence = list(csv.DictReader(f, delimiter="\t"))
    for i, row in enumerate(occurrence):
        code = row["语步"]
        china = float(row["中国出现率(%)"])
        us = float(row["美国出现率(%)"])
        y = 2.95 + i * 0.31
        c.text(1.02, y + 0.01, 0.42, 0.24, code, 8.5, TEXT, bold=True, valign="middle")
        for j, (value, color) in enumerate([(china, BLUE), (us, ORANGE)]):
            by = y + 0.025 + j * 0.125
            width = plot_w * value / 100
            c.rect(plot_x, by, width, 0.085, color, radius=0.018)
            c.text(plot_x + width + 0.06, by - 0.025, 0.48, 0.14,
                   f"{value:.1f}", 7.5, MUTED, valign="middle")

    # Three key gaps.
    c.rect(9.15, 2.18, 3.25, 4.12, PALE_BLUE, radius=0.11)
    c.text(9.39, 2.37, 2.77, 0.38, "关键差距（百分点）", 11, BLUE, bold=True)
    by_code = {row["语步"]: row for row in occurrence}
    key_moves = [
        ("M09", "受众互动与行动引导"),
        ("M08", "合作网络与社会服务"),
        ("M03", "使命愿景与公共价值"),
    ]
    for i, (code, name) in enumerate(key_moves):
        row = by_code[code]
        china = float(row["中国出现率(%)"])
        us = float(row["美国出现率(%)"])
        gap = float(row["美国-中国(百分点)"])
        y = 2.88 + i * 1.00
        c.rect(9.39, y, 2.77, 0.84, WHITE, radius=0.09)
        c.text(9.57, y + 0.10, 2.40, 0.26, f"{code}　{name}", 10, NAVY, bold=True)
        c.text(9.57, y + 0.43, 1.70, 0.27,
               f"{china:.1f}% → {us:.1f}%", 11.5, TEXT, bold=True, valign="middle")
        c.text(11.20, y + 0.40, 0.75, 0.30, f"+{gap:.1f}", 14, ORANGE,
               bold=True, align="right", valign="middle")
    c.text(9.39, 5.94, 2.77, 0.22, "其余7个语步差距均≤10个百分点", 8.5, MUTED, align="center")

    c.text(0.84, 6.40, 8.00, 0.20,
           "注：出现率为该国30个文本中包含相应语步的文本比例。", 8.5, MUTED)
    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "05 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_06(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 4.20, 0.30, "04　研究结果｜语步分布", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 11.20, 0.52, "篇幅与开篇：M05为共同核心，开篇功能不同", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "M05的篇幅占比中位数在两国均约20%；M02更常开启中国文本，M05更常开启美国文本",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    # Left: editable dumbbell chart for median move-length shares.
    c.rect(0.80, 2.18, 7.65, 4.12, WHITE, line=LINE, radius=0.11, width=0.8)
    c.text(1.03, 2.37, 2.90, 0.30, "语步词数占比中位数（%）", 11, BLUE, bold=True)
    c.circle(5.92, 2.43, 0.12, BLUE)
    c.text(6.11, 2.34, 0.55, 0.28, "中国", 9.5, MUTED, valign="middle")
    c.circle(6.86, 2.43, 0.12, ORANGE)
    c.text(7.05, 2.34, 0.55, 0.28, "美国", 9.5, MUTED, valign="middle")

    length_path = ROOT.parent / "analysis" / "table_2_3_move_length.tsv"
    with length_path.open(encoding="utf-8-sig", newline="") as f:
        lengths = list(csv.DictReader(f, delimiter="\t"))
    position_path = ROOT.parent / "analysis" / "table_2_4_move_position_opening.tsv"
    with position_path.open(encoding="utf-8-sig", newline="") as f:
        positions = {row["语步"]: row for row in csv.DictReader(f, delimiter="\t")}

    plot_x, plot_w, scale_max = 1.52, 5.55, 22.0
    # Highlight the shared core move before drawing the grid.
    c.rect(1.00, 4.12, 7.15, 0.28, PALE_BLUE, radius=0.04)
    for tick in [0, 5, 10, 15, 20]:
        x = plot_x + plot_w * tick / scale_max
        c.line(x, 2.88, x, 6.00, LINE, 0.5)
        c.text(x - 0.22, 2.64, 0.44, 0.20, str(tick), 8, MUTED, align="center")

    for i, row in enumerate(lengths):
        code = row["语步"]
        china = float(row["中国占比中位数(%)"])
        us = float(row["美国占比中位数(%)"])
        y = 2.95 + i * 0.30
        x_cn = plot_x + plot_w * china / scale_max
        x_us = plot_x + plot_w * us / scale_max
        c.text(1.02, y + 0.005, 0.40, 0.23, code, 8.5, TEXT, bold=True, valign="middle")
        c.line(x_cn, y + 0.12, x_us, y + 0.12, LINE, 1.2)
        c.circle(x_cn - 0.065, y + 0.055, 0.13, BLUE)
        c.circle(x_us - 0.065, y + 0.055, 0.13, ORANGE)
        if code == "M05":
            c.text(x_us - 0.52, y + 0.025, 0.40, 0.19, f"{us:.1f}", 8, ORANGE,
                   bold=True, align="right", valign="middle")
            c.text(x_cn + 0.12, y + 0.025, 0.40, 0.19, f"{china:.1f}", 8, BLUE,
                   bold=True, valign="middle")
    c.text(1.03, 6.03, 7.05, 0.20, "M05　课程教学与培养目标：在两组文本中均居篇幅占比首位",
           8.5, BLUE, bold=True, align="center")

    # Right: two opening-position comparisons.
    c.rect(8.68, 2.18, 3.72, 4.12, PALE_BLUE, radius=0.11)
    c.text(8.95, 2.37, 3.18, 0.30, "开篇比例（%）", 11, BLUE, bold=True)
    opening_cards = [
        ("M02", "历史沿革与学脉传承", "中国更常以历史开篇"),
        ("M05", "课程教学与培养目标", "美国更常以课程教学开篇"),
    ]
    for i, (code, name, takeaway) in enumerate(opening_cards):
        row = positions[code]
        china = float(row["中国开篇比例(%)"])
        us = float(row["美国开篇比例(%)"])
        y = 2.88 + i * 1.48
        c.rect(8.95, y, 3.18, 1.30, WHITE, radius=0.09)
        c.text(9.12, y + 0.10, 2.84, 0.25, f"{code}　{name}", 10, NAVY, bold=True)
        for j, (label, value, color) in enumerate([("中国", china, BLUE), ("美国", us, ORANGE)]):
            ly = y + 0.42 + j * 0.30
            by = ly + 0.055
            c.text(9.12, ly, 0.38, 0.20, label, 8.5, MUTED, valign="middle")
            c.rect(9.55, by, 1.90 * value / 100, 0.10, color, radius=0.02)
            c.text(9.55 + 1.90 * value / 100 + 0.06, by - 0.035, 0.46, 0.17,
                   f"{value:.1f}", 8, color, bold=True, valign="middle")
        c.text(9.12, y + 1.05, 2.84, 0.19, takeaway, 8.8,
               BLUE if code == "M02" else ORANGE, bold=True, align="center")
    c.text(8.95, 5.88, 3.18, 0.36,
           "开篇＝该语步为文本首个语步\n完整首次位置热图见附录", 8.2, MUTED,
           align="center", spacing=1.15)

    c.text(0.84, 6.40, 7.70, 0.20,
           "注：篇幅占比＝语步词数／文本独立词数，图中取各组中位数。", 8.5, MUTED)
    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "06 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_07(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 4.20, 0.30, "05　研究结果｜语言实现", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 10.80, 0.52, "语言实现：相同语步，不同表达重心", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "关键词差异主要体现在议题选择与语义组合，而非语步是否出现",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    cards = [
        ("M03", "使命愿景与公共价值", "table_2_5_m03_keywords.tsv",
         "人才培养与国际／国家指向", ["talents", "cultivate", "aims", "international"],
         "公共价值与共同体", ["diversity", "equity", "inclusion", "community"]),
        ("M04", "学科价值与研究取向", "table_2_6_m04_keywords.tsv",
         "学科名称与研究范畴", ["foreign", "translation", "linguistics", "literature"],
         "研究主题与兴趣", ["historical", "gender", "race", "art"]),
        ("M05", "课程教学与培养目标", "table_2_7_m05_keywords.tsv",
         "项目供给与能力培养", ["doctoral", "teaching", "talents", "cross-cultural"],
         "学习机会与职业发展", ["minor", "offerings", "seminar", "opportunities", "careers"]),
        ("M08", "合作网络与社会服务", "table_2_8_m08_keywords.tsv",
         "国际交流、合作与培训", ["international", "exchange", "cooperation", "joint", "training"],
         "海外学习与公共参与", ["abroad", "outreach", "community", "engagement", "public"]),
    ]
    positions = [(0.80, 2.18), (6.72, 2.18), (0.80, 4.22), (6.72, 4.22)]
    card_w, card_h = 5.68, 1.82
    for spec, (x, y) in zip(cards, positions):
        code, name, filename, cn_group, cn_terms, us_group, us_terms = spec
        keyword_path = ROOT.parent / "analysis" / filename
        with keyword_path.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f, delimiter="\t"))
        available = {row["词项"] for row in rows}
        assert set(cn_terms + us_terms).issubset(available), (code, cn_terms, us_terms)

        c.rect(x, y, card_w, card_h, WHITE, line=LINE, radius=0.11, width=0.8)
        c.rect(x, y, card_w, 0.06, NAVY)
        c.text(x + 0.20, y + 0.15, 0.55, 0.30, code, 14, BLUE, bold=True, valign="middle")
        c.text(x + 0.86, y + 0.15, 4.58, 0.30, name, 12.5, NAVY, bold=True, valign="middle")
        c.line(x + 0.20, y + 0.51, x + 5.48, y + 0.51, LINE, 0.7)

        for j, (country, group, terms) in enumerate([
            ("中国", cn_group, cn_terms), ("美国", us_group, us_terms)
        ]):
            ry = y + 0.61 + j * 0.61
            is_cn = country == "中国"
            c.rect(x + 0.20, ry, 0.54, 0.26,
                   PALE_BLUE if is_cn else PALE_ORANGE, radius=0.07)
            c.text(x + 0.20, ry, 0.54, 0.26, country, 9,
                   BLUE if is_cn else ORANGE, bold=True, align="center", valign="middle")
            c.text(x + 0.88, ry - 0.01, 4.54, 0.27, group, 10.5, TEXT,
                   bold=True, valign="middle")
            c.text(x + 0.88, ry + 0.29, 4.54, 0.23, " · ".join(terms), 9.2,
                   MUTED, valign="middle")

    c.text(0.84, 6.30, 11.50, 0.20,
           "注：页面列出代表性区分词项；完整每万词频次、log2比率与BH校正q值见附录。", 8.5, MUTED)
    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "07 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def slide_08(deck):
    c = deck.add_slide()
    c.rect(0, 0, 0.16, SLIDE_H, NAVY)
    c.rect(0.16, 0, 0.055, 0.92, ORANGE)

    c.text(0.80, 0.48, 4.20, 0.30, "06　下一阶段计划", 11, BLUE, bold=True)
    c.text(0.80, 0.84, 11.20, 0.52, "下一阶段：解释差异并转化研究成果", 26, NAVY, bold=True)
    c.text(0.80, 1.43, 11.50, 0.42,
           "在定量比较基础上，回答差异为何形成，并将研究发现转化为写作建议与论文",
           14, TEXT, valign="middle")
    c.rect(0.80, 1.94, 0.65, 0.045, ORANGE)

    flow_y, flow_h = 2.30, 3.48
    # Existing quantitative foundation.
    c.rect(0.80, flow_y, 2.55, flow_h, PALE_BLUE, radius=0.12)
    c.rect(1.04, 2.57, 0.78, 0.30, BLUE, radius=0.08)
    c.text(1.04, 2.57, 0.78, 0.30, "已完成", 9, WHITE, bold=True,
           align="center", valign="middle")
    c.text(1.04, 3.02, 2.07, 0.38, "现有定量基础", 15, NAVY, bold=True)
    foundation = ["60个院系文本", "10个核心语步", "语步分布与语言实现"]
    for i, item in enumerate(foundation):
        y = 3.68 + i * 0.48
        c.circle(1.06, y + 0.08, 0.12, BLUE)
        c.text(1.32, y, 1.77, 0.28, item, 10.5, TEXT, valign="middle")
    c.line(1.04, 5.18, 3.10, 5.18, "C9D9E9", 0.8)
    c.text(1.04, 5.36, 2.07, 0.26, "提供可解释的文本证据", 9.5, BLUE, bold=True)

    # Qualitative interpretation.
    c.rect(3.78, flow_y, 3.72, flow_h, WHITE, line=TEAL, radius=0.12, width=1.0)
    c.rect(3.78, flow_y, 3.72, 0.07, TEAL)
    c.text(4.02, 2.58, 3.24, 0.38, "01　质性解释", 15, NAVY, bold=True)
    c.text(4.02, 3.03, 3.24, 0.28, "结合典型文本与语境例证", 10, MUTED)
    qualitative = ["跨文化话语分析", "批评话语分析", "社会文化语境解释"]
    for i, item in enumerate(qualitative):
        y = 3.48 + i * 0.50
        c.rect(4.02, y, 3.24, 0.38, PALE_TEAL, radius=0.08)
        c.text(4.02, y, 3.24, 0.38, item, 10.5, TEAL, bold=True,
               align="center", valign="middle")
    c.text(4.02, 5.14, 3.24, 0.45,
           "重点解释差异背后的\n制度、政策与文化动因", 10, TEXT, align="center", spacing=1.18)

    # Outputs: strategy library and paper.
    c.rect(8.15, flow_y, 4.25, flow_h, WHITE, line=ORANGE, radius=0.12, width=1.0)
    c.rect(8.15, flow_y, 4.25, 0.07, ORANGE)
    c.text(8.40, 2.58, 3.75, 0.38, "02　成果转化", 15, NAVY, bold=True)
    c.rect(8.40, 3.14, 3.75, 0.98, PALE_ORANGE, radius=0.09)
    c.text(8.62, 3.29, 3.31, 0.28, "话语策略建议库", 12, ORANGE, bold=True)
    c.text(8.62, 3.62, 3.31, 0.40,
           "语步组织 · 核心概念 · 国际化\n工具性—人文性目标平衡", 9.5, TEXT, spacing=1.16)
    c.rect(8.40, 4.37, 3.75, 0.98, PALE_BLUE, radius=0.09)
    c.text(8.62, 4.52, 3.31, 0.28, "研究论文", 12, BLUE, bold=True)
    c.text(8.62, 4.85, 3.31, 0.40,
           "整合定量结果与质性解释\n完成论文撰写与修改", 9.5, TEXT, spacing=1.16)

    c.chevron(3.48, 3.79, 0.16, 0.36, TEAL)
    c.chevron(7.82, 3.79, 0.16, 0.36, ORANGE)

    c.rect(0.80, 6.04, 11.60, 0.43, NAVY, radius=0.08)
    c.text(1.02, 6.04, 11.16, 0.43,
           "由“有何差异”推进到“为何形成、如何服务英文简介写作”",
           11, WHITE, bold=True, align="center", valign="middle")

    c.line(0.80, 6.72, 12.55, 6.72, LINE, 0.8)
    c.text(0.80, 6.91, 8.50, 0.28,
           "基于语料库技术的中美外语院系英文简介话语比较分析", 9, MUTED)
    c.text(11.90, 6.91, 0.65, 0.28, "08 / 08", 10, MUTED, align="right")
    deck.save_slide(c)


def main():
    deck = Deck()
    slide_01(deck)
    slide_02(deck)
    slide_03(deck)
    slide_04(deck)
    slide_05(deck)
    slide_06(deck)
    slide_07(deck)
    slide_08(deck)
    deck.save()


if __name__ == "__main__":
    main()
